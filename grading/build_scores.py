"""반별 채점 결과(grades_*.py)를 모아 엑셀 채점표를 만든다.
usage: python3 grading/build_scores.py 통합사회_수행평가_채점표.xlsx
"""
import glob
import importlib.util
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = '맑은 고딕'
P_N, D_N = 7, 6
CIRCLED = '①②③④⑤⑥⑦'
thin = Side(style='thin', color='A6A6A6')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
FILL_P, FILL_D, FILL_T, FILL_IN = 'DDEBF7', 'E2EFDA', 'FFF2CC', 'FFFFE0'


def font(**kw):
    return Font(name=FONT, size=kw.pop('size', 10), **kw)


def fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


def load(path):
    spec = importlib.util.spec_from_file_location(os.path.basename(path)[:-3].replace('-', '_'), path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# 열 배치
C_NO, C_NAME = 1, 2
C_P = 3                       # 청원 ①~⑦: 3~9
C_PCNT, C_PSCORE = 10, 11
C_D = 12                      # 결정 ①~⑥: 12~17
C_DCNT, C_DSCORE = 18, 19
C_SPECIAL, C_TOTAL, C_REASON, C_NOTE = 20, 21, 22, 23
L = get_column_letter


def class_sheet(wb, g):
    ws = wb.create_sheet(g.CLASS)
    # 1~2행: 머리글
    groups = [(C_NO, C_NO, '번호', None), (C_NAME, C_NAME, '이름', None),
              (C_P, C_PSCORE, '국회 전자청원문 (10점)', FILL_P),
              (C_D, C_DSCORE, '헌법재판소 결정문 (10점)', FILL_D),
              (C_SPECIAL, C_SPECIAL, '처리', None), (C_TOTAL, C_TOTAL, '합계\n(20점)', FILL_T),
              (C_REASON, C_REASON, '감점 사유', None), (C_NOTE, C_NOTE, '확인 사항', None)]
    for a, b, text, color in groups:
        if a != b:
            ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        else:
            ws.merge_cells(start_row=1, start_column=a, end_row=2, end_column=a)
        c = ws.cell(1, a, text)
        c.font, c.alignment = font(bold=True), CENTER
        for col in range(a, b + 1):
            for r in (1, 2):
                ws.cell(r, col).fill = fill(color or 'F2F2F2')
                ws.cell(r, col).border = BORDER
    sub = {C_PCNT: '충족', C_PSCORE: '점수', C_DCNT: '충족', C_DSCORE: '점수'}
    sub.update({C_P + i: CIRCLED[i] for i in range(P_N)})
    sub.update({C_D + i: CIRCLED[i] for i in range(D_N)})
    for col, text in sub.items():
        c = ws.cell(2, col, text)
        c.font, c.alignment = font(bold=True), CENTER

    first = 3
    for i, (no, name) in enumerate(g.STUDENTS):
        r = first + i
        pm, dm = g.P.get(no, {}), g.D.get(no, {})
        ws.cell(r, C_NO, no)
        ws.cell(r, C_NAME, name)
        for k in range(P_N):
            ws.cell(r, C_P + k, 'X' if k + 1 in pm else 'O').fill = fill(FILL_IN)
        for k in range(D_N):
            ws.cell(r, C_D + k, 'X' if k + 1 in dm else 'O').fill = fill(FILL_IN)
        ws.cell(r, C_SPECIAL).fill = fill(FILL_IN)
        ws.cell(r, C_PCNT, f'=COUNTIF({L(C_P)}{r}:{L(C_P + P_N - 1)}{r},"O")')
        ws.cell(r, C_PSCORE, f'={L(C_PCNT)}{r}+3')            # 7개 10점 … 0개 3점
        ws.cell(r, C_DCNT, f'=COUNTIF({L(C_D)}{r}:{L(C_D + D_N - 1)}{r},"O")')
        ws.cell(r, C_DSCORE, f'={L(C_DCNT)}{r}+4')            # 6개 10점 … 0개 4점
        ws.cell(r, C_TOTAL, f'=IF({L(C_SPECIAL)}{r}="장기결석",6,{L(C_PSCORE)}{r}+{L(C_DSCORE)}{r})')
        reasons = [pm[k] for k in sorted(pm)] + [dm[k] for k in sorted(dm)]
        ws.cell(r, C_REASON, '\n'.join(reasons))
        ws.cell(r, C_NOTE, g.NOTE.get(no, ''))
        for col in range(1, C_NOTE + 1):
            c = ws.cell(r, col)
            c.border = BORDER
            c.font = font(bold=col == C_TOTAL)
            c.alignment = LEFT if col in (C_REASON, C_NOTE) else CENTER
        if reasons or g.NOTE.get(no):
            ws.row_dimensions[r].height = 15 * max(len(reasons), 1) + 12

    last = first + len(g.STUDENTS) - 1
    avg = last + 1
    ws.cell(avg, C_NAME, '평균').font = font(bold=True)
    for col in (C_PSCORE, C_DSCORE, C_TOTAL):
        c = ws.cell(avg, col, f'=ROUND(AVERAGE({L(col)}{first}:{L(col)}{last}),2)')
        c.font, c.alignment = font(bold=True), CENTER
    for col in range(1, C_NOTE + 1):
        ws.cell(avg, col).border = BORDER
        ws.cell(avg, col).fill = fill('F2F2F2')

    red = PatternFill('solid', start_color='F8CBAD', end_color='F8CBAD')
    ws.conditional_formatting.add(f'{L(C_P)}{first}:{L(C_D + D_N - 1)}{last}',
                                  CellIsRule(operator='equal', formula=['"X"'], fill=red))

    widths = {C_NO: 5, C_NAME: 8, C_PCNT: 5, C_PSCORE: 5, C_DCNT: 5, C_DSCORE: 5,
              C_SPECIAL: 8, C_TOTAL: 7, C_REASON: 70, C_NOTE: 40}
    for col in range(1, C_NOTE + 1):
        ws.column_dimensions[L(col)].width = widths.get(col, 3.6)
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = ws.cell(3, C_P)
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0


def guide_sheet(wb):
    ws = wb.active
    ws.title = '채점 기준'
    rows = [
        ('통합사회2 수행평가 채점표 — 기본권 침해 사례 관련 법률 문서 작성하기 (20점)', True),
        ('', False),
        ('■ 입력 방법', True),
        ('· 반별 시트의 연노랑 칸(기준 ①~⑦, ①~⑥)에 O 또는 X를 입력하면 충족 수·점수·합계가 자동 계산됩니다.', False),
        ('· 장기 미인정 결석 미응시자는 ‘처리’ 칸에 장기결석 이라고 쓰면 합계가 6점이 됩니다.', False),
        ('· 미제출(백지 포함)은 모든 기준을 X로 두면 기본점수(청원문 3점, 결정문 4점)가 됩니다.', False),
        ('', False),
        ('■ 국회 전자청원문 (충족 개수 + 3점: 7개 10점 … 0개 3점)', True),
        ('① 기본권 침해 상황을 구체적으로(누가·언제·어떤 불이익) 서술', False),
        ('② 침해된 기본권의 종류 제시', False),
        ('③ 문제의 원인이 된 법률·제도·정책 제시', False),
        ('④ 관련 헌법 조문을 조·항까지 제시', False),
        ('⑤ 기본권 침해 이유를 근거 자료(기사·통계·사례) 1개 이상 들어 서술', False),
        ('⑥ 국회에 요구하는 구제 내용(청원 취지) 제시', False),
        ('⑦ 시민 참여 방안 1가지 이상 제시', False),
        ('', False),
        ('■ 헌법재판소 결정문 (충족 개수 + 4점: 6개 10점 … 0개 4점)', True),
        ('① 사건 개요를 자신의 말로 정리', False),
        ('② 충돌하는 기본권 또는 헌법적 쟁점', False),
        ('③ 판단 근거가 되는 헌법 조문 제시', False),
        ('④ 합헌 또는 위헌 결론을 명확히 제시', False),
        ('⑤ 공익·사익 등 가치를 고려한 판단 이유', False),
        ('⑥ 결정문 형식(사건 개요–쟁점–판단–결론)', False),
        ('', False),
        ('■ 채점 원칙 (넉넉하게 채점, 근거가 분명할 때만 감점)', True),
        ('· 헌법 조문(청원문 ④, 결정문 ③): 사례에 따라 필요한 경우와 아닌 경우가 있어 감점하지 않습니다(모두 O 처리).', False),
        ('· 근거 자료: 기사·통계·조사·영상·청원 사례 등 출처가 드러나면 인정합니다.', False),
        ('· 시민 참여: 청원 동의, 캠페인, SNS 알리기, 설문, 신고 등 내가 실천할 방법이면 인정합니다.', False),
        ('· 형식·분량: 칸 순서가 조금 달라도 개요–쟁점–판단–결론이 드러나면 인정하고, 분량이 짧은 것만으로는 감점하지 않습니다.', False),
        ('· 결정문 결론: 실제 결정과 달라도 감점하지 않으며, 헌법불합치는 위헌으로 봅니다.', False),
    ]
    for i, (text, bold) in enumerate(rows, 1):
        c = ws.cell(i, 1, text)
        c.font = font(bold=bold, size=12 if i == 1 else 10)
    ws.column_dimensions['A'].width = 110


def main(out):
    wb = Workbook()
    guide_sheet(wb)
    here = os.path.dirname(os.path.abspath(__file__))
    for path in sorted(glob.glob(os.path.join(here, 'grades_*.py')),
                       key=lambda p: [int(x) for x in os.path.basename(p)[7:-3].split('-')]):
        class_sheet(wb, load(path))
    wb.save(out)
    print('saved', out)


if __name__ == '__main__':
    main(sys.argv[1])
